"""
Generate standalone PowerPoint presentations for each question.
Each PPTX has: Slide 1 (question + timer) + one slide per option (diagram + explanation).
"""

import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image


# Colors
NAVY_DARK = RGBColor(0x0D, 0x1B, 0x2A)
NAVY_DARKER = RGBColor(0x09, 0x12, 0x1C)
ACCENT_BLUE = RGBColor(0x1E, 0x88, 0xE5)
TIMER_BG = RGBColor(0x1A, 0x23, 0x7E)
GREEN_ACCENT = RGBColor(0x00, 0xE6, 0x76)
CORRECT_GREEN = RGBColor(0x00, 0xC8, 0x53)
INCORRECT_RED = RGBColor(0xFF, 0x17, 0x44)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xCC, 0xCC, 0xCC)
DARK_TEXT = RGBColor(0x21, 0x21, 0x21)
SLIDE_BG_LIGHT = RGBColor(0xF8, 0xF9, 0xFA)
BADGE_BLUE = RGBColor(0x1E, 0x88, 0xE5)


def create_question_pptx(question: dict, session_num: int, q_num: int):
    """Create a standalone PPTX for a single question."""
    q_dir = Path(f"session{session_num}/exam-strategy/q{q_num}")
    q_dir.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    # Set slide size to 16:9 widescreen (10" x 5.625")
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)

    # Slide 1: Question + Options + Timer
    _create_question_slide(prs, question, session_num, q_num)

    # Slides 2+: One per option
    for letter in sorted(question["options"].keys()):
        _create_option_slide(prs, question, letter, session_num, q_num)

    # Save
    output_path = q_dir / f"q{q_num}_exam_strategy.pptx"
    prs.save(str(output_path))
    print(f"  Created: {output_path}")


def _create_question_slide(prs, question: dict, session_num: int, q_num: int):
    """Create the question slide with options and timer."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Dark background
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = NAVY_DARK

    # Left accent bar
    left_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0),
        Inches(0.05), Inches(5.625)
    )
    left_bar.fill.solid()
    left_bar.fill.fore_color.rgb = ACCENT_BLUE
    left_bar.line.fill.background()

    # Title
    title_box = slide.shapes.add_textbox(
        Inches(0.3), Inches(0.2), Inches(7.0), Inches(0.5)
    )
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f"Session {session_num} - Question {q_num}"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.font.name = "Segoe UI"

    # Question text
    q_box = slide.shapes.add_textbox(
        Inches(0.3), Inches(0.7), Inches(7.2), Inches(1.4)
    )
    tf = q_box.text_frame
    tf.word_wrap = True
    tf.margin_top = Pt(0)
    tf.margin_bottom = Pt(0)
    p = tf.paragraphs[0]
    # Truncate very long questions
    q_text = question["question"]
    if len(q_text) > 350:
        q_text = q_text[:347] + "..."
    p.text = q_text
    p.font.size = Pt(9)
    p.font.color.rgb = WHITE
    p.font.name = "Segoe UI"

    # Options - dynamically sized based on count
    options = question["options"]
    num_options = len(options)
    y_start = 2.2
    available_height = 5.5 - y_start  # Leave small margin at bottom
    spacing = available_height / num_options
    # Truncate option text to fit 2 lines at 8pt in ~6.8" wide box
    # At 8pt Segoe UI, ~6.8" fits roughly 95 chars per line, so 2 lines = ~190
    # But keep it shorter for readability
    max_chars = 90 if num_options >= 6 else 110

    for i, letter in enumerate(sorted(options.keys())):
        opt = options[letter]
        opt_text = opt["text"] if opt["text"] else opt["explanation"][:100]
        # Truncate to fit
        if len(opt_text) > max_chars:
            # Try to cut at a natural boundary
            cut_at = opt_text.rfind('. ', 0, max_chars)
            if cut_at < 40:
                cut_at = opt_text.rfind(': ', 0, max_chars)
            if cut_at < 40:
                cut_at = opt_text.rfind(', ', 0, max_chars)
            if cut_at < 40:
                cut_at = max_chars
            opt_text = opt_text[:cut_at] + "..."

        y = y_start + (i * spacing)
        box_height = spacing - 0.02

        # Badge circle
        badge = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(0.3), Inches(y + 0.05), Inches(0.22), Inches(0.22)
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = BADGE_BLUE
        badge.line.fill.background()

        # Badge letter
        badge_tf = badge.text_frame
        badge_tf.margin_top = Pt(0)
        badge_tf.margin_bottom = Pt(0)
        badge_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        bp = badge_tf.paragraphs[0]
        bp.text = letter
        bp.font.size = Pt(7)
        bp.font.bold = True
        bp.font.color.rgb = WHITE
        bp.alignment = PP_ALIGN.CENTER

        # Option text
        opt_box = slide.shapes.add_textbox(
            Inches(0.6), Inches(y), Inches(7.0), Inches(box_height)
        )
        otf = opt_box.text_frame
        otf.word_wrap = True
        otf.margin_top = Pt(2)
        otf.margin_bottom = Pt(0)
        otf.margin_left = Pt(0)
        op = otf.paragraphs[0]
        op.text = f"{opt_text}"
        op.font.size = Pt(8)
        op.font.color.rgb = LIGHT_GRAY
        op.font.name = "Segoe UI"

    # Timer card (top right)
    timer_card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(7.8), Inches(0.2), Inches(2.0), Inches(2.3)
    )
    timer_card.fill.solid()
    timer_card.fill.fore_color.rgb = TIMER_BG
    timer_card.line.fill.background()

    # Timer label
    timer_label = slide.shapes.add_textbox(
        Inches(7.8), Inches(0.3), Inches(2.0), Inches(0.3)
    )
    tlf = timer_label.text_frame
    tp = tlf.paragraphs[0]
    tp.text = "⏱ TIME"
    tp.font.size = Pt(9)
    tp.font.bold = True
    tp.font.color.rgb = GREEN_ACCENT
    tp.font.name = "Segoe UI"
    tp.alignment = PP_ALIGN.CENTER

    # Embed countdown GIF
    gif_path = Path(f"session{session_num}/exam-strategy/countdown_2min.gif")
    if gif_path.exists():
        slide.shapes.add_picture(
            str(gif_path),
            Inches(7.95), Inches(0.6),
            Inches(1.7), Inches(1.7)
        )


def _create_option_slide(prs, question: dict, letter: str, session_num: int, q_num: int):
    """Create a slide for a single option with diagram and explanation."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    opt = question["options"][letter]
    is_correct = opt["correct"]

    # Light background
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = SLIDE_BG_LIGHT

    # Top banner
    banner = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0),
        Inches(10), Inches(0.9)
    )
    banner.fill.solid()
    banner.fill.fore_color.rgb = NAVY_DARK
    banner.line.fill.background()

    # Status badge
    badge_color = CORRECT_GREEN if is_correct else INCORRECT_RED
    badge_text = "✓ CORRECT" if is_correct else "✗ INCORRECT"

    badge = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.3), Inches(0.2), Inches(1.3), Inches(0.4)
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = badge_color
    badge.line.fill.background()

    badge_tf = badge.text_frame
    badge_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    bp = badge_tf.paragraphs[0]
    bp.text = badge_text
    bp.font.size = Pt(10)
    bp.font.bold = True
    bp.font.color.rgb = WHITE
    bp.alignment = PP_ALIGN.CENTER
    bp.font.name = "Segoe UI"

    # Option text in banner
    opt_text = opt["text"] if opt["text"] else opt["explanation"][:100]
    opt_label = slide.shapes.add_textbox(
        Inches(1.8), Inches(0.2), Inches(7.8), Inches(0.5)
    )
    otf = opt_label.text_frame
    otf.word_wrap = True
    op = otf.paragraphs[0]
    op.text = f"Option {letter}: {opt_text}"
    op.font.size = Pt(10)
    op.font.color.rgb = WHITE
    op.font.name = "Segoe UI"

    # Diagram (left side)
    diagram_path = Path(f"session{session_num}/exam-strategy/q{q_num}/option_{letter.lower()}.png")
    if diagram_path.exists():
        # Get image dimensions for proportional sizing
        with Image.open(diagram_path) as img:
            img_w, img_h = img.size

        # Target area: 5.8" wide, 4.2" tall
        max_w = 5.8
        max_h = 4.2
        aspect = img_w / img_h

        if aspect > max_w / max_h:
            # Width-constrained
            diag_w = max_w
            diag_h = max_w / aspect
        else:
            # Height-constrained
            diag_h = max_h
            diag_w = max_h * aspect

        # Center vertically
        y_offset = 1.0 + (max_h - diag_h) / 2

        slide.shapes.add_picture(
            str(diagram_path),
            Inches(0.2), Inches(y_offset),
            Inches(diag_w), Inches(diag_h)
        )

    # Explanation panel (right side)
    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(6.2), Inches(1.1), Inches(3.5), Inches(4.2)
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = WHITE
    panel.line.color.rgb = RGBColor(0xE0, 0xE0, 0xE0)

    # Accent line on left edge of panel
    accent_line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(6.2), Inches(1.3), Inches(0.05), Inches(3.8)
    )
    accent_line.fill.solid()
    accent_line.fill.fore_color.rgb = badge_color
    accent_line.line.fill.background()

    # Explanation text
    header_text = "Why this is correct:" if is_correct else "Why this is incorrect:"
    explanation = opt["explanation"]

    exp_box = slide.shapes.add_textbox(
        Inches(6.4), Inches(1.3), Inches(3.1), Inches(3.9)
    )
    tf = exp_box.text_frame
    tf.word_wrap = True

    # Header
    p = tf.paragraphs[0]
    p.text = header_text
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = badge_color
    p.font.name = "Segoe UI"
    p.space_after = Pt(6)

    # Bullet points (shorter, more numerous)
    bullets = _split_for_pptx(explanation)
    for bullet in bullets:
        p = tf.add_paragraph()
        p.text = f"• {bullet}"
        p.font.size = Pt(8)
        p.font.color.rgb = DARK_TEXT
        p.font.name = "Segoe UI"
        p.space_after = Pt(4)


def _split_for_pptx(text: str) -> list[str]:
    """Split explanation into short, scannable bullet points for PPTX.
    
    Strategy: split at sentence boundaries, then break overly long sentences
    at natural clause boundaries. Target ~60-90 chars per bullet, up to 8 bullets.
    """
    if not text:
        return ["No explanation provided"]

    import re

    # Split on sentence boundaries (period followed by space + capital letter)
    raw_sentences = re.split(r'(?<=[.!?])\s+(?=[A-Z])', text.strip())
    raw_sentences = [s.strip().rstrip('.').strip() for s in raw_sentences if s.strip()]

    # Break sentences > 90 chars at natural clause boundaries
    fragments = []
    clause_splitters = re.compile(
        r',\s+(?:which |allowing |enabling |ensuring |eliminating |rather than |instead of |as it |it (?:is|does|can|will) )'
        r'|;\s+'
        r'|\s+(?:Additionally|However|Furthermore|This ensures|This eliminates|This enhances|This allows)\b',
        flags=re.IGNORECASE
    )

    for sentence in raw_sentences:
        if len(sentence) > 90:
            parts = clause_splitters.split(sentence)
            parts = [p.strip().rstrip('.').strip() for p in parts if p and p.strip()]
            if len(parts) > 1:
                for part in parts:
                    if part and len(part) > 10:
                        fragments.append(part)
            else:
                # Try splitting at ", " followed by a verb phrase
                alt_parts = re.split(r',\s+(?=\w)', sentence, maxsplit=2)
                if len(alt_parts) > 1 and all(len(p) > 15 for p in alt_parts):
                    for part in alt_parts:
                        part = part.strip().rstrip('.')
                        if part and len(part) > 10:
                            fragments.append(part)
                else:
                    fragments.append(sentence)
        elif len(sentence) > 10:
            fragments.append(sentence)

    # Deduplicate near-identical fragments
    seen = set()
    unique = []
    for f in fragments:
        key = f.lower()[:40]
        if key not in seen:
            seen.add(key)
            unique.append(f)

    # Capitalize first letter of each bullet
    bullets = []
    for f in unique:
        if f:
            bullets.append(f[0].upper() + f[1:] if len(f) > 1 else f.upper())

    # Return up to 8 bullets
    return bullets[:8] if bullets else [text[:150]]


def generate_all_pptx():
    """Generate PPTX files for all sessions."""
    with open("extracted_questions.json") as f:
        sessions = json.load(f)

    for session_num_str, questions in sessions.items():
        session_num = int(session_num_str)
        print(f"\nSession {session_num}: generating {len(questions)} PPTX files")

        for q in questions:
            q_num = q["number"]
            create_question_pptx(q, session_num, q_num)


if __name__ == "__main__":
    generate_all_pptx()
