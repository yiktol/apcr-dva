"""
Extract exam questions from APCR-DVA Answer Key documents.
Handles multiple document formats across sessions 1-5.
"""

import re
import json
from pathlib import Path
from docx import Document


def extract_session1(doc_path: str) -> list[dict]:
    """
    Session 1 format:
    - Questions start with 'N:' pattern
    - "The correct answer is X:" or "The correct answers are:"
    - Correct: "ServiceName (Letter): explanation"
    - Incorrect: same format under "The incorrect options/answers are:"
    """
    doc = Document(doc_path)
    paragraphs = [p.text.strip() for p in doc.paragraphs]

    questions = []
    q_starts = []
    for i, text in enumerate(paragraphs):
        if re.match(r'^\d+:\s+\w', text):
            q_starts.append(i)

    for qi, start_idx in enumerate(q_starts):
        end_idx = q_starts[qi + 1] if qi + 1 < len(q_starts) else len(paragraphs)
        block = paragraphs[start_idx:end_idx]
        q_data = _parse_s1_block(block, qi + 1)
        if q_data:
            questions.append(q_data)

    return questions


def _parse_s1_block(block: list[str], q_num: int) -> dict:
    """Parse a session 1 question block."""
    question_lines = []
    rest_start = 0
    for i, line in enumerate(block):
        if not line:
            continue
        if line.startswith("The main keywords") or line.startswith("The correct answer"):
            rest_start = i
            break
        question_lines.append(line)

    question_text = " ".join(question_lines)
    question_text = re.sub(r'^\d+:\s*', '', question_text)

    correct_letters = set()
    options = {}
    section = None  # "correct" or "incorrect"

    for line in block[rest_start:]:
        if not line:
            continue
        if line.startswith("The main keywords"):
            continue
        if line.startswith("This solution"):
            # Continuation explanation for correct answer
            if options:
                last_correct = [k for k in options if options[k]["correct"]]
                if last_correct:
                    options[last_correct[-1]]["explanation"] += " " + line
            continue

        # Section headers
        if re.search(r'correct answers? (is|are)', line.lower()) and 'incorrect' not in line.lower():
            section = "correct"
            # Inline: "The correct answer is C: Send the JSON..."
            m = re.match(r'.*correct answer is ([A-F]):\s*(.*)', line)
            if m:
                letter = m.group(1)
                correct_letters.add(letter)
                options[letter] = {
                    "text": m.group(2).strip(),
                    "explanation": m.group(2).strip(),
                    "correct": True
                }
            else:
                # "The correct answer is A:" or "The correct answers are A, C, and E:"
                m = re.search(r'correct answers? (?:is|are)\s+([A-F](?:,\s*[A-F])*(?:,?\s*and\s+[A-F])?)', line)
                if m:
                    letters_str = m.group(1)
                    for letter in re.findall(r'[A-F]', letters_str):
                        correct_letters.add(letter)
            continue
        elif re.search(r'incorrect (options|answers) are', line.lower()):
            section = "incorrect"
            continue

        # Parse option: "ServiceName (Letter): explanation"
        m = re.match(r'^(.+?)\s*\(([A-F])\):\s*(.*)', line)
        if m:
            letter = m.group(2)
            is_correct = section == "correct" or letter in correct_letters
            if is_correct:
                correct_letters.add(letter)
            options[letter] = {
                "text": m.group(1).strip(),
                "explanation": m.group(3).strip(),
                "correct": is_correct
            }
            continue

        # Parse: "A) text"
        m = re.match(r'^([A-F])\)\s*(.*)', line)
        if m:
            letter = m.group(1)
            is_correct = section == "correct" or letter in correct_letters
            if is_correct:
                correct_letters.add(letter)
            options[letter] = {
                "text": m.group(2).strip(),
                "explanation": m.group(2).strip(),
                "correct": is_correct
            }
            continue

        # Parse: "A. text" (letter followed by period and space)
        m = re.match(r'^([A-F])\.\s+(.*)', line)
        if m:
            letter = m.group(1)
            is_correct = section == "correct" or letter in correct_letters
            if is_correct:
                correct_letters.add(letter)
            options[letter] = {
                "text": m.group(2).strip(),
                "explanation": m.group(2).strip(),
                "correct": is_correct
            }
            continue

        # Continuation - append to last option in current section
        if section == "correct":
            correct_opts = [k for k in options if options[k]["correct"]]
            if correct_opts:
                options[correct_opts[-1]]["explanation"] += " " + line
        elif section == "incorrect":
            incorrect_opts = [k for k in options if not options[k]["correct"]]
            if incorrect_opts:
                options[incorrect_opts[-1]]["explanation"] += " " + line

    return {
        "number": q_num,
        "question": question_text,
        "options": options,
        "correct_letters": sorted(correct_letters)
    }


def extract_sessions_2_5(doc_path: str, session_num: int) -> list[dict]:
    """
    Sessions 2-5 format varies:
    - Questions: 'Q{N}:', '{N})', '#{N}'
    - Options on Normal lines: "A) text" or "A: Correct/Incorrect. explanation"
    - Explanations on List Paragraph lines or Normal (Web) lines
    - Correct = no "Incorrect" keyword in explanation
    """
    doc = Document(doc_path)
    paragraphs = [(p.text.strip(), p.style.name) for p in doc.paragraphs]

    questions = []
    q_starts = []
    for i, (text, style) in enumerate(paragraphs):
        if not text:
            continue
        if "Answer Key" in text:
            continue
        if re.match(r'^Q\d+\s*:', text):
            q_starts.append(i)
        elif re.match(r'^\d+\)\s', text):
            q_starts.append(i)
        elif re.match(r'^#\d+\s', text):
            q_starts.append(i)

    for qi, start_idx in enumerate(q_starts):
        end_idx = q_starts[qi + 1] if qi + 1 < len(q_starts) else len(paragraphs)
        block = [(paragraphs[j][0], paragraphs[j][1]) for j in range(start_idx, end_idx)]
        q_data = _parse_s2_5_block(block, qi + 1)
        if q_data:
            questions.append(q_data)

    return questions


def _parse_s2_5_block(block: list[tuple], q_num: int) -> dict:
    """Parse a session 2-5 question block."""
    if not block:
        return None

    # Build structured list: [(letter, option_text, explanation_lines)]
    question_lines = []
    entries = []  # list of (letter, option_text, [explanation_lines])
    current_letter = None
    current_option_text = ""
    current_explanations = []

    for text, style in block:
        if not text:
            continue
        # Skip "For more information" lines
        if text.startswith("For more information"):
            continue

        # Check if this is an option line (starts with A-E followed by ) or : or .)
        m = re.match(r'^([A-E])\s*[):.]\s*(.*)', text)
        if m and style in ("Normal", "Normal (Web)"):
            # Save previous entry
            if current_letter:
                entries.append((current_letter, current_option_text, current_explanations))
            elif question_lines:
                pass  # question lines already collected

            current_letter = m.group(1)
            remaining = m.group(2).strip()

            # Check if this line IS the explanation (format: "A: Correct. explanation")
            if re.match(r'^(Correct|Incorrect)', remaining):
                current_option_text = ""
                current_explanations = [remaining]
            else:
                current_option_text = remaining
                current_explanations = []
        elif current_letter is None:
            # Still in question section
            question_lines.append(text)
        else:
            # Explanation line for current option
            current_explanations.append(text)

    # Save last entry
    if current_letter:
        entries.append((current_letter, current_option_text, current_explanations))

    # Clean question text
    question_text = " ".join(question_lines)
    question_text = re.sub(r'^(Q\d+|#\d+|\d+\))\s*[:.]?\s*', '', question_text)

    # Determine correctness for each option
    options = {}
    for letter, option_text, explanations in entries:
        full_explanation = " ".join(explanations)

        # Determine if correct
        is_correct = False
        if re.match(r'^Correct\.?\s', full_explanation):
            is_correct = True
            full_explanation = re.sub(r'^Correct\.?\s*', '', full_explanation)
        elif re.match(r'^Incorrect', full_explanation):
            is_correct = False
            full_explanation = re.sub(r'^Incorrect\.?\s*', '', full_explanation)
        elif 'incorrect' in full_explanation.lower():
            is_correct = False
        elif full_explanation and 'incorrect' not in full_explanation.lower():
            # No "incorrect" keyword and has explanation = likely correct
            # But check for negative indicators
            negative_indicators = [
                'not a good practice', 'not suitable', 'does not',
                'cannot', 'not for use', 'not likely', 'not optimized',
                'very similar to', 'not the best way', 'not necessarily',
                'would not', 'will not'
            ]
            has_negative = any(ind in full_explanation.lower() for ind in negative_indicators)
            if not has_negative:
                is_correct = True

        options[letter] = {
            "text": option_text,
            "explanation": full_explanation,
            "correct": is_correct
        }

    correct_letters = sorted([k for k, v in options.items() if v["correct"]])

    if not options:
        return None

    return {
        "number": q_num,
        "question": question_text,
        "options": options,
        "correct_letters": correct_letters
    }


def extract_all_sessions() -> dict:
    """Extract questions from all 5 sessions."""
    sessions = {}

    s1_path = "ppt/03 - APCR-DVA - Exam Strategy Session 1 - Answer Key.docx"
    sessions[1] = extract_session1(s1_path)

    session_files = {
        2: "ppt/05 - APCR-DVA - Exam Strategy Session 2 - Answer Key.docx",
        3: "ppt/07 - APCR-DVA - Exam Strategy Session 3 - Answer Key.docx",
        4: "ppt/09 - APCR-DVA - Exam Strategy Session 4 - Answer Key.docx",
        5: "ppt/11 - APCR-DVA - Exam Strategy Session 5 - Answer Key.docx",
    }

    for session_num, path in session_files.items():
        sessions[session_num] = extract_sessions_2_5(path, session_num)

    return sessions


if __name__ == "__main__":
    sessions = extract_all_sessions()

    for session_num, questions in sessions.items():
        print(f"\n{'='*60}")
        print(f"Session {session_num}: {len(questions)} questions")
        print(f"{'='*60}")
        for q in questions:
            print(f"\n  Q{q['number']}: {q['question'][:100]}...")
            print(f"  Correct: {q['correct_letters']}")
            for letter in sorted(q['options'].keys()):
                opt = q['options'][letter]
                status = "✓" if opt['correct'] else "✗"
                expl = opt['explanation'][:80]
                print(f"    {status} {letter}: {expl}...")

    # Save to JSON
    output_path = Path("extracted_questions.json")
    output = {}
    for session_num, questions in sessions.items():
        output[str(session_num)] = questions

    with open(output_path, "w") as f:
        json.dump(output, f, indent=2)

    print(f"\n\nSaved to {output_path}")
